import time
import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from modelling.text_preprocessing import TextPreprocessing
from modelling.immense import test, get_or_create_bert_post_embeddings
from os.path import join, exists
from utils import load_from_pickle, aggregate_bert_post_embeddings


def create_results_workbook(losses, conf_nes, confs):
    """Create one sheet per loss and one vertically stacked table per NE pair."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    table_rows = {}
    thin = Side(style="thin", color="BFBFBF")
    edge = Side(style="thin", color="000000")
    headers = ["C", "R", "S", "Prec", "Rec", "F1", "Acc",
               "Prec", "Rec", "F1", "Prec", "Rec", "F1"]
    for loss in losses:
        sheet = workbook.create_sheet(loss)
        sheet.sheet_view.showGridLines = False
        for col in range(1, 14):
            sheet.column_dimensions[get_column_letter(col)].width = 8 if col <= 3 else 11
        for index, (ne_rel, ne_spat) in enumerate(conf_nes):
            top = 1 + index * (len(confs) + 5)
            sheet.merge_cells(start_row=top, start_column=1, end_row=top, end_column=13)
            title = sheet.cell(top, 1, f"Node embeddings: relational = {ne_rel}, spatial = {ne_spat}")
            title.font = Font(bold=True, size=12)
            sheet.row_dimensions[top].height = 24
            for start, end, label in [(1, 3, "Configuration"), (4, 7, "All users"),
                                      (8, 10, "Safe"), (11, 13, "Risky")]:
                sheet.merge_cells(start_row=top + 1, start_column=start,
                                  end_row=top + 1, end_column=end)
                sheet.cell(top + 1, start, label)
            for col, label in enumerate(headers, 1):
                sheet.cell(top + 2, col, label)
            for row in sheet.iter_rows(min_row=top + 1, max_row=top + 2 + len(confs), min_col=1, max_col=13):
                for cell in row:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = Border(left=edge if cell.column in (1, 4, 8, 11) else thin,
                                         right=edge if cell.column in (3, 7, 10, 13) else thin,
                                         top=thin, bottom=thin)
                    if cell.row <= top + 2:
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill("solid", fgColor="E8EEF4")
                    elif cell.column >= 4:
                        cell.number_format = "0.00"
                sheet.row_dimensions[row[0].row].height = 21
            for offset, conf in enumerate(confs):
                row = top + 3 + offset
                table_rows[loss, (ne_rel, ne_spat), conf] = row
                for col, enabled in enumerate(conf, 1):
                    if enabled:
                        sheet.cell(row, col, "\u2713").font = Font(color="008040", bold=True, size=14)
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.print_area = f"A1:M{sheet.max_row}"
    return workbook, table_rows


def write_report(sheet, row, report):
    """Keep full precision in Excel; display metrics to two decimal places."""
    metrics = [report["macro avg"][key] for key in ("precision", "recall", "f1-score")]
    metrics.append(report["accuracy"])
    for label in ("0", "1"):
        metrics.extend(report[label][key] for key in ("precision", "recall", "f1-score"))
    for col, value in enumerate(metrics, 4):
        sheet.cell(row, col, float(value))


def main_test(results_path="test_results_rf_w2v.xlsx"):
    now = time.time()
    with open("parameters.yaml", 'r') as params_file:
        params = yaml.safe_load(params_file)
        dataset_general_params = params["dataset_general_params"]
        test_dataset_params = params["test_dataset_params"]
        model_params = params["model_params"]

    field_name_id = dataset_general_params["field_id"]
    field_name_text = dataset_general_params["field_text"]
    field_label = dataset_general_params["field_label"]
    bert = dataset_general_params["bert"]

    test_df = test_dataset_params["test_df"]
    posts_df_path = dataset_general_params["posts_df"]
    path_rel = test_dataset_params["test_social_net"]
    path_spat = test_dataset_params["test_spatial_net"]
    consider_content = test_dataset_params["consider_content"]
    consider_rel = test_dataset_params["consider_rel"]
    consider_spat = test_dataset_params["consider_spat"]
    separator = test_dataset_params["separator"]

    ne_dim_rel = int(model_params["ne_dim_rel"])
    ne_dim_spat = int(model_params["ne_dim_spat"])
    loss = model_params["loss"]
    word_emb_size = int(model_params["word_emb_size"])
    models_dir = model_params["dir_models"]
    methodology = model_params["methodology"]

    test_df = pd.read_csv(test_df)
    mod_rel = mod_spat = softmax_model = w2v_model = dang_ae = safe_ae = None       # depending on the configuration, some of them will stay None, others won't

    posts_df = pd.read_csv(posts_df_path)

    if bert:
        models_dir += "/with_bert"
        word_emb_size = 768
        softmax_path = join(models_dir, "softmax_model.pkl")

        if exists(softmax_path):
            print(f"[SOFTMAX] Loading softmax model from {softmax_path}")
            softmax_model = load_from_pickle(softmax_path)
            softmax_model.eval()
        else:
            raise FileNotFoundError(f"[ERROR!] Softmax model not found in {softmax_path}")

        posts_embs_dict = get_or_create_bert_post_embeddings(model_dir=models_dir, text_field_name=field_name_text,
                                                             posts_df=posts_df)
        users_embs_dict = aggregate_bert_post_embeddings(users_df=test_df, posts_embs_dict=posts_embs_dict,
                                                         posts_df=posts_df)
    else:
        models_dir += "/with_w2v"
        w2v_path = join(models_dir, "w2v_{}.pkl".format(word_emb_size))
        w2v_model = load_from_pickle(w2v_path)
        dang_ae = load_from_pickle(join(models_dir, "autoencoderdang_{}.pkl".format(word_emb_size)))
        safe_ae = load_from_pickle(join(models_dir, "autoencodersafe_{}.pkl".format(word_emb_size)))

        tok = TextPreprocessing()
        posts = tok.token_dict(test_df, text_field_name=field_name_text, id_field_name=field_name_id)
        users_embs_dict = w2v_model.text_to_vec(posts)

    mod_dir_rel = join(models_dir, "node_embeddings", "rel")
    mod_dir_spat = join(models_dir, "node_embeddings", "spat")

    losses = ["weighted", "focal"]   #

    #(True, False, False),
    confs = [(True, False, True), (True, True, False), (True, True, True), (False, False, True), (False, True, False), (False, True, True)]
    #confs = [(True, True, False), (True, True, True), (False, False, True), (False, True, False), (False, True, True)]
    conf_nes = [(512, 512), (256, 256), (128, 128)] #, (512, 512), (256, 256), (128, 128)
    if bert:
        conf_nes.append((768, 768))
    co = 1

    workbook, table_rows = create_results_workbook(losses, conf_nes, confs)
    for loss in losses:
        for conf_ne in conf_nes:
            ne_dim_rel, ne_dim_spat = conf_ne
            for conf in confs:
                consider_content, consider_rel, consider_spat = conf
                mod_rel = mod_spat = forest_rel = forest_spat = None
                mlp_name = "mlp"
                if consider_content:
                    mlp_name += "_content_{}".format(word_emb_size)
                if consider_rel:
                    mod_rel_name = f"graphsage_{ne_dim_rel}_{word_emb_size}_{loss}.pkl"
                    mod_rel = load_from_pickle(join(mod_dir_rel, mod_rel_name))
                    mlp_name += "_rel_{}".format(ne_dim_rel)
                    if methodology == "random_forest":
                        forest_rel = load_from_pickle(join(mod_dir_rel, "random_forest_{}_{}_{}.pkl".format(ne_dim_rel, word_emb_size, loss)))
                if consider_spat:
                    mod_spat_name = f"graphsage_{ne_dim_spat}_{word_emb_size}_{loss}.pkl"
                    mod_spat = load_from_pickle(join(mod_dir_spat, mod_spat_name))
                    mlp_name += "_spat_{}".format(ne_dim_spat)
                    if methodology == "random_forest":
                        forest_spat = load_from_pickle(join(mod_dir_rel, "random_forest_{}_{}_{}.pkl".format(ne_dim_spat, word_emb_size, loss)))

                if methodology == "random_forest":
                    mlp_name += "_random_forest"
                mlp_name += "_{}.pkl".format(loss)
                mlp = load_from_pickle(join(models_dir, "mlp", mlp_name))
                print(f"\n{co}/{len(losses) * len(confs) * len(conf_nes)}")
                print(mlp_name.upper())

                report = test(df=test_df, users_embeddings_dict=users_embs_dict, ae_risky=dang_ae, ae_safe=safe_ae, mlp=mlp,
                     mod_rel=mod_rel, softmax_model=softmax_model, mod_spat=mod_spat, rel_net_path=path_rel,
                     spat_net_path=path_spat, field_name_id=field_name_id, field_name_label=field_label, bert=bert,
                     consider_content=consider_content, consider_rel=consider_rel, consider_spat=consider_spat,
                     separator=separator, forest_rel=forest_rel, forest_spat=forest_spat)
                write_report(workbook[loss], table_rows[loss, conf_ne, conf], report)
                # Preserve completed results if a later configuration fails.
                workbook.save(results_path)
                print(f"ELAPSED TIME: {time.time()-now}")
                co += 1
    print(f"Classification reports saved to {results_path}")


if __name__ == "__main__":
    main_test()
